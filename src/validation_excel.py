"""月間検証結果のExcel出力.

3シート: 月間検証（1日1行）/ 問題一覧（Issue 1件1行）/ 取込情報。
個人別の勤怠明細は出力しない。DB/Streamlitに依存しない。
"""

from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from src.constants import (
    VALIDATION_STATUS_ERROR,
    VALIDATION_STATUS_OK,
    VALIDATION_STATUS_WARNING,
)
from src.export_excel import (
    TITLE_FONT,
    WEEKDAY_LABELS,
    new_workbook,
    set_cell,
    weekday_fill,
    workbook_to_bytes,
    write_header_row,
    year_month_label,
)
from src.models import MonthlyValidationResult, ValidationIssue
from src.month_utils import weekday_index
from src.validation_display import (
    REQUIREMENT_MISSING_LABEL,
    format_role_cell,
    format_skill_cell,
)

SHEET_DAILY = "月間検証"
SHEET_ISSUES = "問題一覧"
SHEET_IMPORT = "取込情報"

STATUS_FILLS = {
    VALIDATION_STATUS_OK: PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
    VALIDATION_STATUS_WARNING: PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
    VALIDATION_STATUS_ERROR: PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid"),
}
_SEVERITY_ORDER = {VALIDATION_STATUS_ERROR: 0, VALIDATION_STATUS_WARNING: 1}


def validation_excel_filename(year_month: str) -> str:
    return f"cleaning_staff_validation_{year_month}.xlsx"


def _set_widths(ws, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


# ---------------------------------------------------------------------------
# Sheet 1: 月間検証
# ---------------------------------------------------------------------------


def _build_daily_sheet(
    wb: Workbook, result: MonthlyValidationResult, role_names: dict[int, str]
) -> None:
    ws = wb.create_sheet(SHEET_DAILY)
    role_ids = list(role_names)
    headers = ["日付", "曜日", "清掃勤務人数", "必要人数", "最大人数"]
    headers += [role_names[r] for r in role_ids]
    headers += ["スキル条件", "未登録人数", "UNKNOWN人数", "判定"]
    write_header_row(ws, 1, headers)

    for row, day in enumerate(result.days, start=2):
        fill = weekday_fill(day.work_date)
        values = [
            day.work_date,
            WEEKDAY_LABELS[weekday_index(day.work_date)],
            day.actual_cleaning_staff,
            day.required_staff if day.requirement_defined else REQUIREMENT_MISSING_LABEL,
            day.max_total_staff,
        ]
        values += [format_role_cell(day, r) for r in role_ids]
        values += [
            format_skill_cell(day),
            day.unmatched_working_count,
            day.unknown_cleaning_shift_count,
        ]
        for col, value in enumerate(values, start=1):
            set_cell(ws, row, col, value, fill=fill if col <= 2 else None, align_center=True)
        set_cell(
            ws, row, len(values) + 1, day.status,
            bold=True, fill=STATUS_FILLS[day.status], align_center=True,
        )

    ws.freeze_panes = "C2"
    _set_widths(ws, [12, 6, 13, 11, 10] + [16] * len(role_ids) + [20, 11, 13, 11])


# ---------------------------------------------------------------------------
# Sheet 2: 問題一覧
# ---------------------------------------------------------------------------


def sorted_issues(result: MonthlyValidationResult) -> list[ValidationIssue]:
    """ERROR → WARNING の順、同じ重大度内は日付順（同日内は検出順）."""
    issues = [i for day in result.days for i in day.issues]
    return sorted(issues, key=lambda i: (_SEVERITY_ORDER.get(i.severity, 9), i.work_date))


def _build_issues_sheet(wb: Workbook, result: MonthlyValidationResult) -> None:
    ws = wb.create_sheet(SHEET_ISSUES)
    write_header_row(ws, 1, ["日付", "Severity", "Issue Code", "内容", "必要", "確定", "最大可能"])
    for row, issue in enumerate(sorted_issues(result), start=2):
        set_cell(ws, row, 1, issue.work_date, align_center=True)
        set_cell(
            ws, row, 2, issue.severity,
            bold=True, fill=STATUS_FILLS[issue.severity], align_center=True,
        )
        set_cell(ws, row, 3, issue.code)
        set_cell(ws, row, 4, issue.message)
        set_cell(ws, row, 5, issue.required, align_center=True)
        set_cell(ws, row, 6, issue.actual, align_center=True)
        set_cell(ws, row, 7, issue.possible, align_center=True)
    ws.freeze_panes = "A2"
    _set_widths(ws, [12, 10, 28, 90, 8, 8, 10])


# ---------------------------------------------------------------------------
# Sheet 3: 取込情報
# ---------------------------------------------------------------------------


def _build_import_sheet(wb: Workbook, result: MonthlyValidationResult, exported_at: str) -> None:
    ws = wb.create_sheet(SHEET_IMPORT)
    record = result.import_record
    ws.cell(row=1, column=1, value=f"{year_month_label(result.year_month)} 清掃体制検証").font = TITLE_FONT

    items = [
        ("対象月", year_month_label(result.year_month)),
        ("取込ファイル名", record.source_filename),
        ("取込日時", record.imported_at),
        ("従業員数", record.employee_count),
        ("入力済みシフトセル数", record.shift_count),
        ("未登録スタッフ数", record.unmatched_count),
        ("OK日数", result.count_status(VALIDATION_STATUS_OK)),
        ("WARNING日数", result.count_status(VALIDATION_STATUS_WARNING)),
        ("ERROR日数", result.count_status(VALIDATION_STATUS_ERROR)),
        ("要件未設定日数", result.requirement_missing_days),
        ("出力日時", exported_at),
    ]
    write_header_row(ws, 3, ["項目", "値"])
    for row, (label, value) in enumerate(items, start=4):
        set_cell(ws, row, 1, label, bold=True)
        set_cell(ws, row, 2, value)
    _set_widths(ws, [22, 40])


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def build_validation_workbook(
    result: MonthlyValidationResult,
    role_names: dict[int, str],
    exported_at: str | None = None,
) -> Workbook:
    """月間検証結果の3シートを持つワークブックを作る. ACTIVE取込のない結果は ValueError."""
    if not result.has_import:
        raise ValueError(f"{result.year_month} の取込データがないため出力できません。")
    exported_at = exported_at or datetime.now().isoformat(timespec="seconds")
    wb = new_workbook()
    _build_daily_sheet(wb, result, role_names)
    _build_issues_sheet(wb, result)
    _build_import_sheet(wb, result, exported_at)
    return wb


def export_validation_excel(
    result: MonthlyValidationResult,
    role_names: dict[int, str],
    exported_at: str | None = None,
) -> bytes:
    """st.download_button 向けに .xlsx の bytes を返す."""
    return workbook_to_bytes(build_validation_workbook(result, role_names, exported_at))
