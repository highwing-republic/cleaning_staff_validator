"""Excel出力の共通基盤.

検証結果シート（日別サマリー・問題一覧）は Phase 2 で本モジュールの上に実装する。
DB/Streamlitに依存しない。
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from src.month_utils import parse_year_month, weekday_index

WEEKDAY_LABELS = ("月", "火", "水", "木", "金", "土", "日")

HEADER_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
SATURDAY_FILL = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
SUNDAY_FILL = PatternFill(start_color="F2DCDB", end_color="F2DCDB", fill_type="solid")
TITLE_FONT = Font(bold=True, size=14)
THIN_BORDER = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
CENTER = Alignment(horizontal="center", vertical="center")


def year_month_label(year_month: str) -> str:
    year, month = parse_year_month(year_month)
    return f"{year}年{month}月"


def weekday_fill(work_date: str) -> PatternFill | None:
    """土曜・日曜の網掛け. 平日はNone."""
    wd = weekday_index(work_date)
    if wd == 5:
        return SATURDAY_FILL
    if wd == 6:
        return SUNDAY_FILL
    return None


def set_cell(ws: Worksheet, row: int, col: int, value, *, bold: bool = False,
             fill: PatternFill | None = None, align_center: bool = False,
             number_format: str | None = None) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.border = THIN_BORDER
    if bold:
        cell.font = Font(bold=True)
    if fill is not None:
        cell.fill = fill
    if align_center:
        cell.alignment = CENTER
    if number_format is not None:
        cell.number_format = number_format


def write_header_row(ws: Worksheet, row: int, headers: list[str]) -> None:
    for col, header in enumerate(headers, start=1):
        set_cell(ws, row, col, header, bold=True, fill=HEADER_FILL, align_center=True)


def new_workbook() -> Workbook:
    """既定の空シートを削除したワークブックを作る（シートは呼び出し側で順に追加する）."""
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def workbook_to_bytes(wb: Workbook) -> bytes:
    """st.download_button向けにワークブックをbytesとして返す."""
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
