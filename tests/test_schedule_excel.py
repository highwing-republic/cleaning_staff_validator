"""勤務表Excel（印刷・配布用）のテスト（Phase 11）.

Phase 4の検証Excel（validation_excel）とは別物で、こちらは勤務表そのものを出力する。
"""

import pytest
from openpyxl import load_workbook

from src import repositories as repo
from src import services
from src.database import get_connection, initialize_database
from src.models import DailyRequirementInput, RoleRequirementInput
from src.period_utils import period_dates
from src.schedule_excel import (
    APP_NAME,
    REMARKS_HEADER,
    SHEET_DAILY,
    SHEET_INFO,
    SHEET_ISSUES,
    SHEET_SCHEDULE,
    build_schedule_workbook,
    schedule_excel_filename,
)
from src.schedule_output_display import DRAFT_NOTICE
from src.staffing_validation import STAFF_SHORTAGE

LEADER, CHECKER, CLEANER = 1, 2, 3
EVERY_WEEKDAY = [0, 1, 2, 3, 4, 5, 6]

DATES = period_dates("2026-10-20", 3)
DATES_14 = period_dates("2026-10-20", 14)
DAY = DATES[0]


@pytest.fixture()
def conn():
    c = get_connection(":memory:")
    initialize_database(c)
    yield c
    c.close()


def add_staff(conn, code, name, *, role_id=CLEANER, skill_level=3, end_time="15:30"):
    return repo.create_staff(
        conn, code, name, role_id, skill_level, "清掃",
        standard_start_time="09:00", standard_end_time=end_time,
        weekdays=EVERY_WEEKDAY,
    )


def seed(conn, dates=DATES, required_total_staff=2, role_requirements=(), **kwargs):
    ids = {
        "leader": add_staff(conn, "0101", "リーダー田中", role_id=LEADER, skill_level=5),
        "cleaner": add_staff(conn, "0102", "清掃Aさん", role_id=CLEANER, skill_level=3),
        "part": add_staff(conn, "0103", "短時間Bさん", role_id=CLEANER, skill_level=2,
                          end_time="13:00"),
    }
    repo.save_period_requirements(
        conn,
        list(dates),
        [
            DailyRequirementInput(
                work_date=d, required_total_staff=required_total_staff, **kwargs
            )
            for d in dates
        ],
        [
            RoleRequirementInput(d, role_id, count)
            for d in dates
            for role_id, count in role_requirements
        ],
    )
    result = services.generate_schedule(conn, list(dates))
    run_id, errors = services.save_generated_schedule(conn, list(dates), result)
    assert errors == []
    return ids


def workbook(conn, dates=DATES):
    return build_schedule_workbook(
        services.build_schedule_export(conn, list(dates), exported_at="2026-10-19T08:00:00")
    )


def row_values(ws, row):
    return [ws.cell(row=row, column=col).value for col in range(1, ws.max_column + 1)]


def column_values(ws, column):
    return [ws.cell(row=row, column=column).value for row in range(1, ws.max_row + 1)]


def staff_row(ws, name):
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == name:
            return row_values(ws, row)
    raise AssertionError(f"スタッフ行が見つかりません: {name}")


def sheet_text(ws):
    return [
        str(cell.value)
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    ]


# ---------------------------------------------------------------------------
# シート構成
# ---------------------------------------------------------------------------


def test_workbook_has_the_four_sheets(conn):
    seed(conn)
    wb = workbook(conn)
    assert wb.sheetnames == [SHEET_SCHEDULE, SHEET_DAILY, SHEET_ISSUES, SHEET_INFO]


def test_filename_contains_the_period(conn):
    assert schedule_excel_filename(DATES) == "cleaning_schedule_2026-10-20_2026-10-22.xlsx"
    assert schedule_excel_filename([]) == "cleaning_schedule.xlsx"


def test_excel_bytes_can_be_reopened(conn):
    seed(conn)
    data = services.export_schedule_excel(conn, DATES)
    assert data[:2] == b"PK"
    import io

    reopened = load_workbook(io.BytesIO(data))
    assert reopened.sheetnames[0] == SHEET_SCHEDULE


def test_export_does_not_save_anything_to_the_database(conn):
    seed(conn)
    before = [
        tuple(r)
        for r in conn.execute(
            "SELECT work_date, staff_id, is_working FROM schedule_assignments ORDER BY 1, 2"
        )
    ]
    services.export_schedule_excel(conn, DATES)
    after = [
        tuple(r)
        for r in conn.execute(
            "SELECT work_date, staff_id, is_working FROM schedule_assignments ORDER BY 1, 2"
        )
    ]
    assert after == before
    assert conn.execute("SELECT count(*) FROM schedule_runs").fetchone()[0] == 1


# ---------------------------------------------------------------------------
# 勤務表シート
# ---------------------------------------------------------------------------


def test_schedule_sheet_lists_staff_and_work_times(conn):
    seed(conn, required_total_staff=3)
    ws = workbook(conn)[SHEET_SCHEDULE]

    assert ws.cell(row=1, column=1).value == APP_NAME
    names = column_values(ws, 1)
    assert "リーダー田中" in names
    assert "清掃Aさん" in names
    assert "短時間Bさん" in names

    leader = staff_row(ws, "リーダー田中")
    assert leader[1] == "リーダー"
    assert "09:00-15:30" in leader
    assert "09:00-13:00" in staff_row(ws, "短時間Bさん")


def test_schedule_sheet_marks_days_off_with_a_single_character(conn):
    ids = seed(conn, required_total_staff=3)
    services.save_manual_schedule_changes(conn, ids["leader"], {DAY: (False, True)})
    ws = workbook(conn)[SHEET_SCHEDULE]
    assert "休" in staff_row(ws, "リーダー田中")


def test_schedule_sheet_does_not_show_lock_or_source_marks(conn):
    """§39・§40: 固定・手修正の印は印刷用の勤務表へ出さない."""
    ids = seed(conn, required_total_staff=3)
    services.save_manual_schedule_changes(conn, ids["leader"], {DAY: (False, True)})
    text = " ".join(sheet_text(workbook(conn)[SHEET_SCHEDULE]))
    assert "🔒" not in text
    assert "※下書き" in text or "MANUAL" not in text
    assert "MANUAL" not in text
    assert "GENERATED" not in text


def test_schedule_sheet_shows_the_day_status_as_text(conn):
    """§41・§49: 確定・下書き・未作成を色ではなく文字で示す."""
    seed(conn, dates=DATES[:2])
    services.finalize_schedule_day(conn, DAY)
    ws = workbook(conn)[SHEET_SCHEDULE]

    status_row = next(
        row for row in range(1, ws.max_row + 1) if ws.cell(row=row, column=1).value == "状態"
    )
    statuses = row_values(ws, status_row)[2:]
    assert "確定" in statuses
    assert "下書き" in statuses
    assert "未作成" in statuses


def test_schedule_sheet_warns_when_the_period_contains_a_draft(conn):
    """§42: 下書きの日を含む場合は勤務表の上部に明記する."""
    seed(conn)
    assert DRAFT_NOTICE in " ".join(sheet_text(workbook(conn)[SHEET_SCHEDULE]))


def test_schedule_sheet_has_no_draft_notice_when_everything_is_finalized(conn):
    seed(conn)
    for work_date in DATES:
        services.finalize_schedule_day(conn, work_date)
    assert DRAFT_NOTICE not in " ".join(sheet_text(workbook(conn)[SHEET_SCHEDULE]))


def test_schedule_sheet_marks_missing_days(conn):
    """§60: 未作成日は空欄ではなく「未作成」と分かるようにする."""
    seed(conn, dates=DATES[:1])
    ws = workbook(conn)[SHEET_SCHEDULE]
    assert "未作成" in staff_row(ws, "リーダー田中")


def test_schedule_sheet_has_a_remarks_column_for_handwriting(conn):
    """§51: 当日の変更を手書きするための空欄を用意する."""
    seed(conn)
    ws = workbook(conn)[SHEET_SCHEDULE]
    header_row = next(
        row for row in range(1, ws.max_row + 1) if ws.cell(row=row, column=1).value == "スタッフ"
    )
    headers = row_values(ws, header_row)
    assert headers[-1] == REMARKS_HEADER

    remarks_column = len(headers)
    leader_row = next(
        row for row in range(1, ws.max_row + 1)
        if ws.cell(row=row, column=1).value == "リーダー田中"
    )
    assert ws.cell(row=leader_row, column=remarks_column).value is None
    assert ws.row_dimensions[leader_row].height >= 20


def test_schedule_sheet_is_set_up_for_a4_landscape_printing(conn):
    """§46・§47: A4横・横1ページに収める・先頭列とヘッダー固定."""
    seed(conn, dates=DATES_14)
    ws = workbook(conn, DATES_14)[SHEET_SCHEDULE]

    assert ws.page_setup.orientation == "landscape"
    assert str(ws.page_setup.paperSize) == str(ws.PAPERSIZE_A4)
    assert ws.page_setup.fitToWidth == 1
    assert ws.page_setup.fitToHeight == 0
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True
    assert ws.freeze_panes.startswith("C")


def test_schedule_sheet_handles_fourteen_days(conn):
    seed(conn, dates=DATES_14, required_total_staff=2)
    ws = workbook(conn, DATES_14)[SHEET_SCHEDULE]
    header_row = next(
        row for row in range(1, ws.max_row + 1) if ws.cell(row=row, column=1).value == "スタッフ"
    )
    headers = row_values(ws, header_row)
    # スタッフ + 役割 + 14日 + 出勤日数 + 備考
    assert len(headers) == 2 + 14 + 2


def test_schedule_sheet_row_order_is_stable(conn):
    """§36: 役割→従業員番号の順で、毎回同じ並びになる."""
    seed(conn)
    first = column_values(workbook(conn)[SHEET_SCHEDULE], 1)
    second = column_values(workbook(conn)[SHEET_SCHEDULE], 1)
    assert first == second

    names = [n for n in first if n in ("リーダー田中", "清掃Aさん", "短時間Bさん")]
    assert names == ["リーダー田中", "清掃Aさん", "短時間Bさん"]


def test_schedule_sheet_keeps_inactive_staff_that_are_in_the_schedule(conn):
    """§37: 無効になったスタッフでも勤務表に行があれば残す."""
    ids = seed(conn, required_total_staff=3)
    repo.deactivate_staff(conn, ids["part"])
    assert "短時間Bさん" in column_values(workbook(conn)[SHEET_SCHEDULE], 1)


def test_schedule_sheet_omits_inactive_staff_without_assignments(conn):
    seed(conn)
    gone = add_staff(conn, "0109", "退職Cさん")
    repo.deactivate_staff(conn, gone)
    assert "退職Cさん" not in column_values(workbook(conn)[SHEET_SCHEDULE], 1)


def test_schedule_sheet_does_not_expose_solver_internals(conn):
    """§66・§67・§68: Solverの内部値・target_days・希望休の破りを勤務表へ出さない."""
    seed(conn)
    text = " ".join(sheet_text(workbook(conn)[SHEET_SCHEDULE]))
    for word in ("target", "deviation", "scaled", "objective", "PREFER_OFF", "希望休"):
        assert word not in text


# ---------------------------------------------------------------------------
# 日別確認シート
# ---------------------------------------------------------------------------


def test_daily_sheet_shows_the_numbers_per_day(conn):
    seed(conn, required_total_staff=2, reserved_rooms=8, role_requirements=((LEADER, 1),))
    services.finalize_schedule_day(conn, DAY)
    ws = workbook(conn)[SHEET_DAILY]

    headers = row_values(ws, 1)
    for label in ("日付", "状態", "予約室数", "必要人数", "配置人数", "人数", "スキル条件"):
        assert label in headers

    first = row_values(ws, 2)
    assert first[0] == DAY
    assert first[1] == "確定"
    assert first[headers.index("予約室数")] == "8"
    assert first[headers.index("必要人数")] == "2"
    assert first[headers.index("配置人数")] == "2"
    assert first[headers.index("人数")] == "OK"


def test_daily_sheet_shows_a_shortage(conn):
    ids = seed(conn, required_total_staff=3)
    services.save_manual_schedule_changes(conn, ids["leader"], {DAY: (False, True)})
    ws = workbook(conn)[SHEET_DAILY]
    headers = row_values(ws, 1)
    assert row_values(ws, 2)[headers.index("人数")] == "1名不足"


def test_daily_sheet_marks_missing_days(conn):
    seed(conn, dates=DATES[:1])
    ws = workbook(conn)[SHEET_DAILY]
    statuses = [row_values(ws, row)[1] for row in range(2, ws.max_row + 1)]
    assert statuses == ["下書き", "未作成", "未作成"]


# ---------------------------------------------------------------------------
# 問題一覧シート
# ---------------------------------------------------------------------------


def test_issues_sheet_lists_one_row_per_issue(conn):
    ids = seed(conn, required_total_staff=3)
    services.save_manual_schedule_changes(conn, ids["leader"], {DAY: (False, True)})
    ws = workbook(conn)[SHEET_ISSUES]

    assert row_values(ws, 1) == ["日付", "Severity", "Issue Code", "内容"]
    codes = [ws.cell(row=row, column=3).value for row in range(2, ws.max_row + 1)]
    assert STAFF_SHORTAGE in codes
    # ERROR が先頭に来る
    assert ws.cell(row=2, column=2).value == "ERROR"


def test_issues_sheet_says_so_when_there_is_nothing_wrong(conn):
    seed(conn, required_total_staff=2)
    for work_date in DATES:
        services.finalize_schedule_day(conn, work_date)
    ws = workbook(conn)[SHEET_ISSUES]
    assert "問題は見つかりませんでした。" in " ".join(sheet_text(ws))


def test_excel_is_produced_even_when_the_schedule_has_problems(conn):
    """§31: 不足があってもExcel出力を禁止しない."""
    seed(conn, required_total_staff=5)
    data = services.export_schedule_excel(conn, DATES)
    assert data[:2] == b"PK"


def test_excel_is_produced_for_a_period_without_any_schedule(conn):
    seed(conn, dates=DATES[:1])
    data = services.export_schedule_excel(conn, period_dates("2026-12-01", 3))
    assert data[:2] == b"PK"


# ---------------------------------------------------------------------------
# 出力情報シート
# ---------------------------------------------------------------------------


def test_info_sheet_reports_the_period_and_counts(conn):
    seed(conn, dates=DATES[:2])
    services.finalize_schedule_day(conn, DAY)
    ws = workbook(conn)[SHEET_INFO]
    items = {
        ws.cell(row=row, column=1).value: ws.cell(row=row, column=2).value
        for row in range(1, ws.max_row + 1)
    }
    assert items["アプリ名"] == APP_NAME
    assert items["対象期間"] == f"{DATES[0]} 〜 {DATES[-1]}"
    assert items["確定日数"] == 1
    assert items["下書き日数"] == 1
    assert items["未作成日数"] == 1
    assert items["出力日時"] == "2026-10-19T08:00:00"


def test_info_sheet_does_not_contain_a_local_path(conn):
    """§45: 個人PCのパス等は出さない."""
    seed(conn)
    text = " ".join(sheet_text(workbook(conn)[SHEET_INFO]))
    assert "C:\\" not in text
    assert "/home/" not in text
    assert "Users" not in text


# ---------------------------------------------------------------------------
# 既存のExcel出力を壊さない（§85）
# ---------------------------------------------------------------------------


def test_attendance_validation_excel_is_untouched():
    from src import validation_excel

    assert validation_excel.SHEET_DAILY == "月間検証"
    assert hasattr(validation_excel, "export_validation_excel")
    assert validation_excel.validation_excel_filename("2026-10") == (
        "cleaning_staff_validation_2026-10.xlsx"
    )


# ---------------------------------------------------------------------------
# 性能（§95）
# ---------------------------------------------------------------------------


def test_excel_generation_is_fast_for_forty_staff_over_fourteen_days(conn):
    import time

    for index in range(1, 41):
        add_staff(conn, f"{index:04d}", f"S{index}")
    repo.save_period_requirements(
        conn, DATES_14,
        [DailyRequirementInput(d, 12) for d in DATES_14], [],
    )
    result = services.generate_schedule(conn, DATES_14)
    services.save_generated_schedule(conn, DATES_14, result)

    started = time.perf_counter()
    data = services.export_schedule_excel(conn, DATES_14)
    elapsed = time.perf_counter() - started

    assert data[:2] == b"PK"
    assert elapsed < 5.0
