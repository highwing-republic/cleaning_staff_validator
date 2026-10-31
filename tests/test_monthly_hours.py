"""月ごとの合計勤務時間（スタッフ単位）のテスト."""

from datetime import date
from pathlib import Path

import pytest
from streamlit.testing.v1 import AppTest

from src import repositories as repo
from src.database import get_connection, initialize_database
from src.monthly_hours import (
    MonthTotal,
    WorkedShift,
    format_hours,
    format_month_total,
    month_dates,
    monthly_totals,
    months_of,
    totals_by_month,
)
from src.period_utils import period_dates
from src.ui_common import DB_PATH_ENV

APP = str(Path(__file__).resolve().parents[1] / "app.py")


# ---------------------------------------------------------------------------
# 計算・表示
# ---------------------------------------------------------------------------


def test_months_of_a_period_across_months():
    assert months_of(period_dates("2026-10-31", 14)) == ["2026-10", "2026-11"]
    assert months_of(period_dates("2026-11-14", 14)) == ["2026-11"]


def test_month_dates():
    assert len(month_dates("2026-11")) == 30
    assert month_dates("2026-02")[-1] == "2026-02-28"


@pytest.mark.parametrize(
    ("minutes", "expected"),
    [(390, "6.5時間"), (405, "6.75時間"), (600, "10時間"), (0, "0時間"), (6630, "110.5時間")],
)
def test_format_hours(minutes, expected):
    assert format_hours(minutes) == expected


def test_format_month_total():
    assert format_month_total(MonthTotal(17, 6630)) == "110.5時間（17日）"
    assert format_month_total(None) == "0時間"


def test_totals_split_by_month():
    shifts = [
        WorkedShift(1, "2026-10-31", "09:00", "15:30"),
        WorkedShift(1, "2026-11-01", "09:00", "13:00"),
        WorkedShift(1, "2026-11-02", "06:00", "18:00"),
        WorkedShift(2, "2026-11-02", "09:00", "15:30"),
    ]
    totals = totals_by_month(shifts)
    assert totals[1]["2026-10"] == MonthTotal(1, 390)
    assert totals[1]["2026-11"] == MonthTotal(2, 240 + 720)
    assert totals[2]["2026-11"] == MonthTotal(1, 390)


# ---------------------------------------------------------------------------
# 保存済みの勤務と合わせる（DB）
# ---------------------------------------------------------------------------


@pytest.fixture
def conn(tmp_path, monkeypatch):
    path = tmp_path / "hours.sqlite"
    monkeypatch.setenv(DB_PATH_ENV, str(path))
    connection = get_connection(path)
    initialize_database(connection)
    yield connection
    connection.close()


def add_staff(conn, code, start="09:00", end="15:30"):
    return repo.create_staff(conn, code, f"S{code}", 3, 3, standard_start_time=start,
                             standard_end_time=end, weekdays=list(range(7)))


def save_day(conn, work_date, staff_times):
    """勤務表を保存した状態を作る. staff_times = {staff_id: (start, end) or None(休み)}."""
    with conn:
        conn.execute("INSERT INTO schedule_days (work_date, status, updated_at) VALUES (?, 'DRAFT', '2026-11-01T00:00:00')",
                     (work_date,))
        for sid, times in staff_times.items():
            conn.execute(
                "INSERT INTO schedule_assignments (work_date, staff_id, is_working, start_time, end_time, locked, source, updated_at) "
                "VALUES (?, ?, ?, ?, ?, 0, 'GENERATED', '2026-11-01T00:00:00')",
                (work_date, sid, int(times is not None), *(times or (None, None))),
            )


def test_new_draft_adds_to_hours_already_saved_this_month(conn):
    """11/1〜13 に保存済みの勤務があれば、11/14〜 の新しい勤務案と合わせて11月の合計にする."""
    sid = add_staff(conn, "1")
    for day in ("2026-11-02", "2026-11-03", "2026-11-04"):
        save_day(conn, day, {sid: ("09:00", "15:30")})
    save_day(conn, "2026-11-05", {sid: None})  # 休みは数えない
    period = period_dates("2026-11-14", 14)
    draft = [WorkedShift(sid, "2026-11-16", "09:00", "13:00"), WorkedShift(sid, "2026-11-17", "09:00", "15:30")]

    totals = monthly_totals(conn, period, draft)
    assert totals[sid]["2026-11"] == MonthTotal(5, 390 * 3 + 240 + 390)


def test_draft_replaces_saved_days_inside_the_period(conn):
    """期間内の日は、保存済みではなく新しい勤務案で数える（作り直したときに二重に数えない）."""
    sid = add_staff(conn, "1")
    save_day(conn, "2026-11-16", {sid: ("09:00", "15:30")})
    period = period_dates("2026-11-14", 14)
    totals = monthly_totals(conn, period, [WorkedShift(sid, "2026-11-17", "09:00", "13:00")])
    assert totals[sid]["2026-11"] == MonthTotal(1, 240)


def test_saved_schedule_counts_the_whole_month(conn):
    sid = add_staff(conn, "1")
    save_day(conn, "2026-10-30", {sid: ("09:00", "15:30")})
    save_day(conn, "2026-11-20", {sid: ("09:00", "15:30")})
    totals = monthly_totals(conn, period_dates("2026-10-31", 14))
    assert totals[sid]["2026-10"] == MonthTotal(1, 390)
    assert totals[sid]["2026-11"] == MonthTotal(1, 390)


# ---------------------------------------------------------------------------
# かんたん作成の③
# ---------------------------------------------------------------------------


def press(at, label):
    next(b for b in at.button if b.label == label).click().run()
    assert not at.exception


def test_step3_shows_month_totals_next_to_the_name(conn):
    sid = add_staff(conn, "1")
    add_staff(conn, "2")
    save_day(conn, "2026-11-02", {sid: ("09:00", "15:30")})  # 前の期間に保存済みの1日
    at = AppTest.from_file(APP, default_timeout=60).run()
    press(at, "勤務表をつくる")
    at.date_input[0].set_value(date(2026, 11, 14))
    press(at, "この期間で始める")
    for d in period_dates("2026-11-14", 14):
        at.number_input(key=f"gw_count_{d}").set_value(1)
    at.checkbox(key="gw_checked_preferences").set_value(True)
    press(at, "人数を保存して次へ")
    press(at, "勤務表をつくる")

    table = at.dataframe[0].value
    assert list(table.columns[:2]) == ["スタッフ", "11月の合計"]
    row = table[table["スタッフ"] == "S1"].iloc[0]
    worked = int(row["勤務日数"].rstrip("日"))
    assert row["11月の合計"] == f"{format_hours(390 * (worked + 1))}（{worked + 1}日）"
    assert row["勤務時間"] == format_hours(390 * worked)
    assert any("休憩は差し引いていません" in c.value for c in at.caption)


# ---------------------------------------------------------------------------
# 印刷・Excel・⑤ 勤務表調整
# ---------------------------------------------------------------------------


def saved_month(conn):
    """10/30 と 11/2〜3 に勤務を保存し、期間 10/31〜11/13 で見る."""
    sid = add_staff(conn, "1")
    other = add_staff(conn, "2")
    save_day(conn, "2026-10-30", {sid: ("09:00", "15:30"), other: None})  # 期間の外（10月）
    save_day(conn, "2026-11-02", {sid: ("09:00", "15:30"), other: ("09:00", "13:00")})
    save_day(conn, "2026-11-03", {sid: ("09:00", "13:00"), other: None})
    return sid, other, period_dates("2026-10-31", 14)


def test_schedule_export_carries_month_totals(conn):
    from src import services

    sid, other, period = saved_month(conn)
    export = services.build_schedule_export(conn, period)
    assert export.months == ["2026-10", "2026-11"]
    rows = {row.staff.staff_id: row for row in export.staff_rows}
    assert rows[sid].period_minutes == 390 + 240
    assert rows[sid].month_totals == {"2026-10": MonthTotal(1, 390), "2026-11": MonthTotal(2, 630)}
    assert rows[other].month_totals == {"2026-11": MonthTotal(1, 240)}


def test_print_shows_month_totals(conn):
    from src import services
    from src.schedule_print import build_schedule_print_html

    sid, _, period = saved_month(conn)
    html = build_schedule_print_html(services.build_schedule_export(conn, period))
    assert '<th scope="col" class="month-total">10月の合計</th>' in html
    assert '<th scope="col" class="month-total">11月の合計</th>' in html
    assert '<span class="shift-time">10.5時間</span><span class="shift-time">2日</span>' in html
    assert "休憩は差し引いていません" in html


def test_excel_shows_hours_and_month_totals(conn):
    from io import BytesIO

    from openpyxl import load_workbook

    from src import services
    from src.schedule_excel import SHEET_SCHEDULE

    saved_month(conn)
    period = period_dates("2026-10-31", 14)
    ws = load_workbook(BytesIO(services.export_schedule_excel(conn, period)))[SHEET_SCHEDULE]
    header_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == "スタッフ")
    headers = [c.value for c in ws[header_row]]
    row = next(r for r in range(header_row + 1, ws.max_row + 1) if ws.cell(row=r, column=1).value == "S1")
    values = dict(zip(headers, [c.value for c in ws[row]]))
    assert values["勤務時間"] == "10.5時間"
    assert values["10月の合計"] == "6.5時間（1日）"
    assert values["11月の合計"] == "10.5時間（2日）"
    assert headers[-1] == "当日変更・備考"


def test_schedule_page_grid_shows_month_totals(conn):
    saved_month(conn)
    page = str(Path(__file__).resolve().parents[1] / "pages" / "07_schedule.py")
    at = AppTest.from_file(page, default_timeout=60)
    at.session_state["schedule_period_start"] = date(2026, 10, 31)
    at.session_state["schedule_period_days"] = 14
    at.run()
    assert not at.exception
    grid = at.dataframe[0].value
    assert list(grid.columns[:3]) == ["スタッフ", "10月の合計", "11月の合計"]
    row = grid[grid["スタッフ"] == "S1"].iloc[0]
    assert row["10月の合計"] == "6.5時間（1日）"
    assert row["11月の合計"] == "10.5時間（2日）"
    assert row["勤務時間"] == "10.5時間"
