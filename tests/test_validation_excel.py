"""月間検証結果Excel（validation_excel）と表示フォーマッタのテスト."""

import io

import pytest
from openpyxl import load_workbook

from src.models import (
    AttendanceImportRecord,
    AttendanceShiftInput,
    DailyRequirementInput,
    MonthlyValidationResult,
    RoleRequirementInput,
    StaffInput,
)
from src.month_utils import get_month_dates
from src.staffing_validation import validate_month_staffing
from src.validation_display import format_count_cell, format_role_cell, format_skill_cell
from src.validation_excel import (
    SHEET_DAILY,
    SHEET_IMPORT,
    SHEET_ISSUES,
    build_validation_workbook,
    export_validation_excel,
    sorted_issues,
    validation_excel_filename,
)

YM = "2026-09"
DATES = get_month_dates(YM)
LEADER, CHECKER, CLEANER = 1, 2, 3
ROLE_NAMES = {LEADER: "リーダー", CHECKER: "チェッカー", CLEANER: "クリーナー"}
STAFF = [
    StaffInput(1, "0001", "L", LEADER, 5, "清掃"),
    StaffInput(2, "0002", "C", CHECKER, 4, "清掃"),
    StaffInput(3, "0003", "W", CLEANER, 2, "清掃"),
]
RECORD = AttendanceImportRecord(
    7, YM, "shift_2026-09.csv", "2026-09-26T10:00:00", 4, 95, 1, "ACTIVE"
)


def _cleaning(code, work_date, staff_id=None):
    return AttendanceShiftInput(
        employee_code=code, work_date=work_date, raw_shift="09:00-15:30",
        shift_type="TIME_RANGE", available_for_cleaning=True, staff_id=staff_id,
        department="清掃", start_minutes=540, end_minutes=930,
    )


def _result():
    """9/1 他: OK / 9/2: ERROR(人数不足) / 9/3: WARNING(未登録) / 9/4: 要件未設定 / 9/5: スキル条件なし."""
    shifts = [_cleaning(s.employee_code, d, s.staff_id) for d in DATES for s in STAFF]
    shifts.append(_cleaning("0099", DATES[2]))
    reqs = []
    for d in DATES:
        if d == DATES[3]:
            continue
        required = 4 if d == DATES[1] else 3
        level, count = (None, 0) if d == DATES[4] else (4, 2)
        reqs.append(DailyRequirementInput(d, required, 5, None, None, level, count))
    roles = [RoleRequirementInput(d, r, 1) for d in DATES for r in (LEADER, CHECKER)]
    days = validate_month_staffing(YM, shifts, reqs, roles, STAFF, ROLE_NAMES)
    return MonthlyValidationResult(YM, RECORD, days)


def _workbook(result=None):
    data = export_validation_excel(result or _result(), ROLE_NAMES, exported_at="2026-09-26T12:00:00")
    return load_workbook(io.BytesIO(data))


def _rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


# ---------------------------------------------------------------------------
# 表示フォーマッタ
# ---------------------------------------------------------------------------


def test_format_count_cell():
    assert format_count_cell(1, 1, 1) == "1 / 1"
    assert format_count_cell(2, 1, 3) == "2 / 1"
    assert format_count_cell(0, 1, 1) == "0 / 1（最大1）"


def test_format_role_and_skill_cells():
    days = _result().days
    ok, missing, no_skill = days[0], days[3], days[4]
    assert format_role_cell(ok, LEADER) == "1 / 1"
    assert format_role_cell(ok, CLEANER) == "1 / 0"
    assert format_role_cell(missing, LEADER) == "1 / -"
    assert format_skill_cell(ok) == "Lv4+ 2 / 2"
    assert format_skill_cell(no_skill) == "-"
    assert format_skill_cell(missing) == "-"


# ---------------------------------------------------------------------------
# ワークブック
# ---------------------------------------------------------------------------


def test_filename_contains_month():
    assert validation_excel_filename("2026-09") == "cleaning_staff_validation_2026-09.xlsx"


def test_three_sheets_in_order():
    assert _workbook().sheetnames == [SHEET_DAILY, SHEET_ISSUES, SHEET_IMPORT]
    assert [SHEET_DAILY, SHEET_ISSUES, SHEET_IMPORT] == ["月間検証", "問題一覧", "取込情報"]


def test_no_import_is_rejected():
    with pytest.raises(ValueError):
        build_validation_workbook(MonthlyValidationResult(YM), ROLE_NAMES)


@pytest.mark.parametrize("year_month,days", [("2026-09", 30), ("2026-10", 31), ("2028-02", 29)])
def test_daily_sheet_has_one_row_per_day(year_month, days):
    dates = get_month_dates(year_month)
    result = MonthlyValidationResult(
        year_month,
        AttendanceImportRecord(1, year_month, "a.csv", "t", 0, 0, 0, "ACTIVE"),
        validate_month_staffing(year_month, [], [], [], [], ROLE_NAMES),
    )
    rows = _rows(_workbook(result)[SHEET_DAILY])
    assert len(rows) == 1 + days
    assert [r[0] for r in rows[1:]] == dates


def test_daily_sheet_headers_have_dynamic_role_columns():
    header = _rows(_workbook()[SHEET_DAILY])[0]
    assert header == [
        "日付", "曜日", "清掃勤務人数", "必要人数", "最大人数",
        "リーダー", "チェッカー", "クリーナー",
        "スキル条件", "未登録人数", "UNKNOWN人数", "判定",
    ]
    custom = {**ROLE_NAMES, 4: "インスペクター"}
    data = export_validation_excel(_result(), custom)
    assert "インスペクター" in _rows(load_workbook(io.BytesIO(data))[SHEET_DAILY])[0]


def test_daily_sheet_values_and_statuses():
    ws = _workbook()[SHEET_DAILY]
    rows = {r[0]: r for r in _rows(ws)[1:]}

    ok = rows[DATES[0]]
    assert ok == [DATES[0], "火", 3, 3, 5, "1 / 1", "1 / 1", "1 / 0", "Lv4+ 2 / 2", 0, 0, "OK"]
    assert rows[DATES[1]][3] == 4
    assert rows[DATES[1]][-1] == "ERROR"
    assert rows[DATES[2]][2] == 4  # 未登録も清掃勤務人数に含む
    assert rows[DATES[2]][9] == 1  # 未登録人数
    assert rows[DATES[2]][-1] == "WARNING"
    # 要件未設定
    assert rows[DATES[3]][3] == "要件未設定"
    assert rows[DATES[3]][4] is None
    assert rows[DATES[3]][5] == "1 / -"
    assert rows[DATES[3]][8] == "-"
    assert rows[DATES[3]][-1] == "WARNING"
    # スキル条件なし
    assert rows[DATES[4]][8] == "-"
    assert rows[DATES[4]][-1] == "OK"


def test_status_cells_are_visually_distinguished():
    ws = _workbook()[SHEET_DAILY]
    status_col = 12
    colors = {
        ws.cell(row, status_col).value: ws.cell(row, status_col).fill.start_color.rgb[-6:]
        for row in range(2, ws.max_row + 1)
    }
    assert len({colors["OK"], colors["WARNING"], colors["ERROR"]}) == 3


def test_issue_sheet_lists_error_then_warning_by_date():
    rows = _rows(_workbook()[SHEET_ISSUES])
    assert rows[0] == ["日付", "Severity", "Issue Code", "内容", "必要", "確定", "最大可能"]
    body = rows[1:]
    assert [(r[0], r[1], r[2]) for r in body] == [
        (DATES[1], "ERROR", "STAFF_SHORTAGE"),
        (DATES[2], "WARNING", "UNMATCHED_STAFF"),
        (DATES[3], "WARNING", "REQUIREMENT_MISSING"),
    ]
    shortage = body[0]
    assert (shortage[4], shortage[5], shortage[6]) == (4, 3, 3)
    assert "不足" in shortage[3]
    # 値のない項目は空欄
    assert body[2][4:] == [None, None, None]


def test_ok_days_not_in_issue_sheet():
    dates_in_issues = {r[0] for r in _rows(_workbook()[SHEET_ISSUES])[1:]}
    assert DATES[0] not in dates_in_issues
    assert DATES[4] not in dates_in_issues


def test_sorted_issues_error_first():
    severities = [i.severity for i in sorted_issues(_result())]
    assert severities == sorted(severities, key=lambda s: 0 if s == "ERROR" else 1)


def test_import_sheet_values():
    rows = _rows(_workbook()[SHEET_IMPORT])
    assert rows[0][0] == "2026年9月 清掃体制検証"
    items = {r[0]: r[1] for r in rows[3:]}
    assert items == {
        "対象月": "2026年9月",
        "取込ファイル名": "shift_2026-09.csv",
        "取込日時": "2026-09-26T10:00:00",
        "従業員数": 4,
        "入力済みシフトセル数": 95,
        "未登録スタッフ数": 1,
        "OK日数": 27,
        "WARNING日数": 2,
        "ERROR日数": 1,
        "要件未設定日数": 1,
        "出力日時": "2026-09-26T12:00:00",
    }


def test_workbook_does_not_contain_personal_shift_details():
    wb = _workbook()
    text = " ".join(
        str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if v
    )
    assert "09:00-15:30" not in text
    assert "0099" not in text
