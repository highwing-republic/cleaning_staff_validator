import io

from openpyxl import load_workbook

from src.export_excel import (
    HEADER_FILL,
    SATURDAY_FILL,
    SUNDAY_FILL,
    new_workbook,
    set_cell,
    weekday_fill,
    workbook_to_bytes,
    write_header_row,
    year_month_label,
)


def test_year_month_label():
    assert year_month_label("2026-09") == "2026年9月"


def test_weekday_fill():
    # 2026-10-01 は木曜、10/3 が土曜、10/4 が日曜
    assert weekday_fill("2026-10-01") is None
    assert weekday_fill("2026-10-03") is SATURDAY_FILL
    assert weekday_fill("2026-10-04") is SUNDAY_FILL


def test_new_workbook_has_no_sheets():
    assert new_workbook().sheetnames == []


def test_workbook_round_trip_with_header_and_cells():
    wb = new_workbook()
    ws = wb.create_sheet("日別サマリー")
    write_header_row(ws, 1, ["日付", "判定"])
    set_cell(ws, 2, 1, "2026-10-01", align_center=True)
    set_cell(ws, 2, 2, 1.5, bold=True, number_format="0.0")

    loaded = load_workbook(io.BytesIO(workbook_to_bytes(wb)))
    assert loaded.sheetnames == ["日別サマリー"]
    sheet = loaded["日別サマリー"]
    assert [sheet.cell(1, c).value for c in (1, 2)] == ["日付", "判定"]
    assert sheet.cell(1, 1).font.bold
    assert sheet.cell(1, 1).fill.start_color.rgb.endswith(HEADER_FILL.start_color.rgb[-6:])
    assert sheet.cell(2, 1).value == "2026-10-01"
    assert sheet.cell(2, 1).alignment.horizontal == "center"
    assert sheet.cell(2, 2).font.bold
    assert sheet.cell(2, 2).number_format == "0.0"
