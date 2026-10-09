"""MVPの通し検証（Phase 12）.

実運用に近い12名・14日のデータで、
  スタッフ設定 → 勤務希望 → 予約・必要人数 → 勤務案作成 → 下書き保存
  → 手修正・固定 → 再生成 → 日別確定 → 勤務表検証 → Excel出力
まで一周することを確認する。

シナリオは module スコープで1回だけ実行し、各段階の観測値を記録しておく。
テストは記録を見るだけなので、実行順に依存しない。
実DB（data/app.db）は触らず、pytest の一時ディレクトリだけを使う。
"""


import pytest
from openpyxl import load_workbook

from src import repositories as repo
from src import services
from src.constants import (
    SPECIAL_SKILL_HEAVY_WORK,
    VALIDATION_STATUS_ERROR,
)
from src.database import get_connection, initialize_database
from src.models import (
    DailyRequirementInput,
    RoleRequirementInput,
    StaffDatePreferenceInput,
)
from src.period_utils import period_dates
from src.schedule_excel import SHEET_DAILY, SHEET_ISSUES, SHEET_SCHEDULE
from src.schedule_validation_adapter import SCHEDULE_DRAFT
from src.staffing_validation import STAFF_SHORTAGE

LEADER, CHECKER, CLEANER = 1, 2, 3

PERIOD_START = "2026-10-20"  # 火曜
PERIOD_DAYS = 14
DATES = period_dates(PERIOD_START, PERIOD_DAYS)

MON, TUE, WED, THU, FRI, SAT, SUN = range(7)
WEEKDAYS_FULL = (MON, TUE, WED, THU, FRI, SAT, SUN)
WEEKDAYS_WEEKEND = (THU, FRI, SAT, SUN, MON)
WEEKDAYS_FIVE = (MON, TUE, WED, THU, FRI)
WEEKDAYS_FOUR = (MON, TUE, WED, THU)
WEEKDAYS_THREE = (TUE, THU, SAT)

# (従業員番号, 氏名, ロール, 総合スキル, 開始, 終了, 通常勤務曜日, 目標/週, 最大連勤, 力仕事可)
STAFF_PLAN = [
    ("1001", "若林リーダー", LEADER, 5, "09:00", "15:30", WEEKDAYS_FIVE, 5, 5, True),
    ("1002", "大野リーダー", LEADER, 4, "09:00", "15:30", WEEKDAYS_WEEKEND, 4, 4, False),
    ("2001", "森チェッカー", CHECKER, 5, "09:00", "15:30", WEEKDAYS_FULL, None, 6, False),
    ("2002", "岡チェッカー", CHECKER, 4, "09:30", "14:30", WEEKDAYS_THREE, 3, 3, False),
    ("3001", "阿部さん", CLEANER, 4, "09:00", "15:30", WEEKDAYS_FULL, 5, 5, True),
    ("3002", "井上さん", CLEANER, 3, "09:00", "15:30", WEEKDAYS_FIVE, 4, 4, False),
    ("3003", "上田さん", CLEANER, 3, "09:00", "13:00", WEEKDAYS_FOUR, 4, None, False),
    ("3004", "江口さん", CLEANER, 2, "09:00", "13:00", WEEKDAYS_THREE, 3, None, False),
    ("3005", "小川さん", CLEANER, 5, "09:30", "14:30", WEEKDAYS_WEEKEND, None, 5, True),
    ("3006", "加藤さん", CLEANER, 3, "09:00", "15:30", WEEKDAYS_FULL, 5, 6, False),
    ("3007", "木村さん", CLEANER, 2, "09:00", "13:00", WEEKDAYS_FIVE, 3, None, False),
    ("3008", "久保さん", CLEANER, 4, "09:00", "15:30", WEEKDAYS_FOUR, None, 4, False),
]

# 予約室数と必要人数（閑散日・通常日・繁忙日・清掃不要日を混ぜる）
DEMAND_PLAN = [
    #  (予約室数, 必要人数, 最大人数)
    (12, 5, None),
    (8, 4, None),
    (0, 0, None),      # 予約0室の日（清掃不要）
    (18, 7, None),
    (20, 8, None),     # 繁忙日
    (6, 3, None),
    (4, 2, None),
    (10, 5, 8),
    (14, 6, None),
    (16, 7, None),
    (2, 1, None),
    (0, 0, None),
    (9, 4, None),
    (11, 5, None),
]

ROLE_REQUIREMENT_DAYS = (0, 3, 4, 8, 9)   # Leader 1 / Checker 1 が要る日
SKILL_REQUIREMENT_DAYS = {3: 2, 4: 2, 9: 1}  # index -> スキル4以上の必要人数

ABSOLUTE_OFF_INDEX = 1        # 阿部さんの絶対休み
PREFER_OFF_INDEXES = (6, 7)   # 井上さんのできれば休み
EARLY_LEAVE_INDEX = 2         # 加藤さんの早上がり
LATE_START_INDEX = 3          # 加藤さんの遅出
EXTRA_INDEX = 5               # 日曜。上田さんが通常休みだが勤務可能
COMBO_INDEX = 10              # 絶対休み + できれば休み の併用（仕様どおり許可）
SHORTAGE_INDEX = 12           # 意図的に1名不足にする日


def _staff_id(ids, employee_code):
    return ids[employee_code]


@pytest.fixture(scope="module")
def scenario(tmp_path_factory):
    """MVPの一連の流れを1度だけ実行し、各段階の観測値を返す."""
    db_path = tmp_path_factory.mktemp("mvp_e2e") / "e2e.db"
    conn = get_connection(str(db_path))
    initialize_database(conn)
    observed = {"db_path": db_path}

    # --- 1. スタッフ設定 -----------------------------------------------
    heavy_work_id = next(
        s.special_skill_id
        for s in repo.list_special_skills(conn)
        if s.skill_code == SPECIAL_SKILL_HEAVY_WORK
    )
    ids = {}
    for (
        code, name, role_id, skill, start, end, weekdays, target, consecutive, heavy
    ) in STAFF_PLAN:
        ids[code] = repo.create_staff(
            conn, code, name, role_id, skill, "清掃",
            standard_start_time=start,
            standard_end_time=end,
            weekdays=list(weekdays),
            target_days_per_week=target,
            max_consecutive_days=consecutive,
            special_skill_ids=[heavy_work_id] if heavy else [],
        )
    observed["ids"] = ids
    observed["heavy_work_id"] = heavy_work_id
    observed["staff_details"] = repo.list_staff_details(conn, include_inactive=False)

    # --- 2. 勤務希望（紙から転記する想定） ------------------------------
    absolute_off_staff = _staff_id(ids, "3001")
    prefer_off_staff = _staff_id(ids, "3002")
    time_override_staff = _staff_id(ids, "3006")
    extra_staff = _staff_id(ids, "3003")
    combo_staff = _staff_id(ids, "2002")

    repo.save_staff_period_preferences(
        conn, absolute_off_staff, DATES,
        [StaffDatePreferenceInput(
            absolute_off_staff, DATES[ABSOLUTE_OFF_INDEX], absolute_off=True,
            note="通院のため",
        )],
    )
    repo.save_staff_period_preferences(
        conn, prefer_off_staff, DATES,
        [
            StaffDatePreferenceInput(prefer_off_staff, DATES[i], prefer_off=True)
            for i in PREFER_OFF_INDEXES
        ],
    )
    repo.save_staff_period_preferences(
        conn, time_override_staff, DATES,
        [
            StaffDatePreferenceInput(
                time_override_staff, DATES[EARLY_LEAVE_INDEX], override_end_time="13:00"
            ),
            StaffDatePreferenceInput(
                time_override_staff, DATES[LATE_START_INDEX], override_start_time="11:00"
            ),
        ],
    )
    repo.save_staff_period_preferences(
        conn, extra_staff, DATES,
        [StaffDatePreferenceInput(extra_staff, DATES[EXTRA_INDEX], available_extra=True)],
    )
    # 絶対休み + できれば休み の併用（矛盾として扱わない仕様）
    repo.save_staff_period_preferences(
        conn, combo_staff, DATES,
        [StaffDatePreferenceInput(
            combo_staff, DATES[COMBO_INDEX], absolute_off=True, prefer_off=True
        )],
    )
    observed["preference_staff"] = {
        "absolute_off": absolute_off_staff,
        "prefer_off": prefer_off_staff,
        "time_override": time_override_staff,
        "extra": extra_staff,
        "combo": combo_staff,
    }
    observed["preferences"] = repo.list_preferences_in_period(conn, DATES[0], DATES[-1])
    observed["conditions"] = {
        staff.staff_id: staff.day_conditions
        for staff in services.build_generation_request(conn, DATES).staff
    }

    # --- 3. 予約室数・必要人数 ------------------------------------------
    available_per_day = {
        work_date: sum(
            1
            for conditions in observed["conditions"].values()
            if conditions[work_date].can_work
            and conditions[work_date].time_status == "OK"
        )
        for work_date in DATES
    }
    observed["available_per_day"] = available_per_day

    requirements = []
    for index, (rooms, required, max_total) in enumerate(DEMAND_PLAN):
        work_date = DATES[index]
        available = available_per_day[work_date]
        if index == SHORTAGE_INDEX:
            # 勤務できる人数より1名多く求める日を作る（不足でも案が返ることの確認用）
            required = available + 1
        elif index == EXTRA_INDEX:
            # 通常休みの人に出てもらわないと埋まらない人数にする
            required = available
        else:
            # 繁忙日でも、その日に勤務できる人数を超えては求めない
            # （不足日は上の1日だけにして、他の日の判定を安定させる）
            required = min(required, available)
        requirements.append(
            DailyRequirementInput(
                work_date=work_date,
                required_total_staff=required,
                max_total_staff=max_total,
                reserved_rooms=rooms,
                required_skill_level=4 if index in SKILL_REQUIREMENT_DAYS else None,
                required_skill_count=SKILL_REQUIREMENT_DAYS.get(index, 0),
            )
        )
    role_requirements = [
        RoleRequirementInput(DATES[index], role_id, 1)
        for index in ROLE_REQUIREMENT_DAYS
        for role_id in (LEADER, CHECKER)
    ]
    repo.save_period_requirements(conn, DATES, requirements, role_requirements)
    observed["requirements"] = repo.list_daily_requirements(conn, DATES[0], DATES[-1])
    observed["role_requirements"] = repo.list_role_requirements(conn, DATES[0], DATES[-1])

    # --- 4. 勤務案の作成 -------------------------------------------------
    result = services.generate_schedule(conn, DATES)
    observed["result"] = result

    # --- 5. 下書き保存 ---------------------------------------------------
    run_id, save_errors = services.save_generated_schedule(conn, DATES, result)
    observed["run_id"] = run_id
    observed["save_errors"] = save_errors
    observed["saved"] = services.get_current_schedule(conn, DATES)

    # --- 6. 手修正と固定 -------------------------------------------------
    edit_target = _staff_id(ids, "3006")
    saved_days = observed["saved"]
    working_dates = [
        d for d in DATES
        if (a := saved_days.day(d).assignments.get(edit_target)) and a.is_working
    ]
    off_dates = [
        d for d in DATES
        if (a := saved_days.day(d).assignments.get(edit_target)) and not a.is_working
    ]
    to_off = working_dates[0]
    to_work = next(
        d for d in off_dates
        if observed["conditions"][edit_target][d].can_work
        and observed["conditions"][edit_target][d].time_status == "OK"
    )
    unlock_target = working_dates[1]

    manual_errors = services.save_manual_schedule_changes(
        conn, edit_target,
        {to_off: (False, True), to_work: (True, True), unlock_target: (True, True)},
    )
    observed["manual_errors"] = manual_errors
    observed["manual"] = {
        "staff_id": edit_target,
        "to_off": to_off,
        "to_work": to_work,
        "unlock_target": unlock_target,
    }
    observed["after_manual"] = services.get_current_schedule(conn, DATES)

    # 固定を外せること（Phase 10の仕様どおり、明示操作でのみ外れる）
    services.save_manual_schedule_changes(conn, edit_target, {unlock_target: (True, False)})
    observed["after_unlock"] = services.get_current_schedule(conn, DATES)

    # --- 7. 固定を守って再生成 --------------------------------------------
    preview = services.preview_regenerated_schedule(conn, DATES)
    observed["preview"] = preview
    observed["before_regeneration"] = services.get_current_schedule(conn, DATES)
    apply_run_id, apply_errors = services.apply_regenerated_schedule(conn, DATES, preview)
    observed["apply_run_id"] = apply_run_id
    observed["apply_errors"] = apply_errors
    observed["after_regeneration"] = services.get_current_schedule(conn, DATES)

    # --- 8. 日別の確定 ----------------------------------------------------
    finalized = [DATES[0], DATES[1]]
    for work_date in finalized:
        assert services.finalize_schedule_day(conn, work_date) == []
    observed["finalized_dates"] = finalized

    # --- 9. 勤務表の検証 --------------------------------------------------
    observed["validation"] = services.validate_current_schedule(conn, DATES)

    # --- 10. Excel出力 ----------------------------------------------------
    excel_bytes = services.export_schedule_excel(conn, DATES)
    observed["excel_bytes"] = excel_bytes
    excel_path = db_path.parent / "cleaning_schedule_e2e.xlsx"
    excel_path.write_bytes(excel_bytes)
    observed["excel_path"] = excel_path
    observed["workbook"] = load_workbook(excel_path)

    conn.close()
    return observed


def _worked_dates(result, staff_id):
    return {a.work_date for a in result.working_assignments(staff_id)}


def _assignment(schedule, work_date, staff_id):
    return schedule.day(work_date).assignments.get(staff_id)


# ---------------------------------------------------------------------------
# 1. スタッフ設定
# ---------------------------------------------------------------------------


def test_e2e_setup_master(scenario):
    details = scenario["staff_details"]
    assert len(details) == 12
    assert sum(1 for d in details if d.staff.role_id == LEADER) == 2
    assert sum(1 for d in details if d.staff.role_id == CHECKER) == 2
    assert sum(1 for d in details if d.staff.role_id == CLEANER) == 8


def test_e2e_setup_weekday_patterns(scenario):
    weekday_counts = {len(d.weekdays) for d in scenario["staff_details"]}
    assert {3, 4, 5, 7} <= weekday_counts


def test_e2e_setup_work_times(scenario):
    times = {
        (d.staff.standard_start_time, d.staff.standard_end_time)
        for d in scenario["staff_details"]
    }
    assert ("09:00", "15:30") in times
    assert ("09:00", "13:00") in times
    assert ("09:30", "14:30") in times
    assert all(t != (None, None) for t in times)


def test_e2e_setup_skill_levels(scenario):
    levels = {d.staff.skill_level for d in scenario["staff_details"]}
    assert levels == {2, 3, 4, 5}


def test_e2e_setup_special_skills(scenario):
    heavy = [
        d for d in scenario["staff_details"]
        if scenario["heavy_work_id"] in d.special_skill_ids
    ]
    assert len(heavy) == 3


def test_e2e_setup_target_days(scenario):
    targets = {d.staff.target_days_per_week for d in scenario["staff_details"]}
    assert targets == {None, 3, 4, 5}


# ---------------------------------------------------------------------------
# 2. 勤務希望
# ---------------------------------------------------------------------------


def test_e2e_preferences_cover_every_kind(scenario):
    preferences = scenario["preferences"]
    assert any(p.absolute_off for p in preferences)
    assert any(p.prefer_off for p in preferences)
    assert any(p.available_extra for p in preferences)
    assert any(p.override_end_time for p in preferences)
    assert any(p.override_start_time for p in preferences)
    assert any(p.note for p in preferences)


def test_e2e_absolute_off_blocks_the_day(scenario):
    staff_id = scenario["preference_staff"]["absolute_off"]
    condition = scenario["conditions"][staff_id][DATES[ABSOLUTE_OFF_INDEX]]
    assert condition.can_work is False


def test_e2e_absolute_off_with_prefer_off_is_allowed(scenario):
    """ABSOLUTE_OFF + PREFER_OFF は併用できる仕様（矛盾扱いしない）."""
    staff_id = scenario["preference_staff"]["combo"]
    saved = [
        p for p in scenario["preferences"]
        if p.staff_id == staff_id and p.work_date == DATES[COMBO_INDEX]
    ]
    assert len(saved) == 1
    assert (saved[0].absolute_off, saved[0].prefer_off) == (True, True)
    assert scenario["conditions"][staff_id][DATES[COMBO_INDEX]].can_work is False


def test_e2e_early_leave_and_late_start_change_the_effective_time(scenario):
    staff_id = scenario["preference_staff"]["time_override"]
    early = scenario["conditions"][staff_id][DATES[EARLY_LEAVE_INDEX]]
    late = scenario["conditions"][staff_id][DATES[LATE_START_INDEX]]
    assert early.effective_end_time == "13:00"
    assert late.effective_start_time == "11:00"


def test_e2e_available_extra_opens_a_normally_off_day(scenario):
    staff_id = scenario["preference_staff"]["extra"]
    assert scenario["conditions"][staff_id][DATES[EXTRA_INDEX]].can_work is True


# ---------------------------------------------------------------------------
# 3. 予約・必要人数
# ---------------------------------------------------------------------------


def test_e2e_requirements_cover_the_whole_period(scenario):
    assert len(scenario["requirements"]) == PERIOD_DAYS
    rooms = [r.reserved_rooms for r in scenario["requirements"]]
    assert 0 in rooms
    assert max(rooms) >= 18


def test_e2e_requirements_include_quiet_and_busy_days(scenario):
    required = [r.required_total_staff for r in scenario["requirements"]]
    assert min(required) == 0
    assert max(required) >= 7


def test_e2e_role_requirements_are_saved(scenario):
    by_date = {}
    for r in scenario["role_requirements"]:
        by_date.setdefault(r.work_date, {})[r.role_id] = r.required_count
    for index in ROLE_REQUIREMENT_DAYS:
        assert by_date[DATES[index]] == {LEADER: 1, CHECKER: 1}


def test_e2e_skill_requirements_are_saved(scenario):
    by_date = {r.work_date: r for r in scenario["requirements"]}
    for index, count in SKILL_REQUIREMENT_DAYS.items():
        requirement = by_date[DATES[index]]
        assert (requirement.required_skill_level, requirement.required_skill_count) == (
            4, count
        )


# ---------------------------------------------------------------------------
# 4. 勤務案の作成
# ---------------------------------------------------------------------------


def test_e2e_generate_returns_a_plan(scenario):
    result = scenario["result"]
    assert result.has_solution
    assert len(result.days) == PERIOD_DAYS
    assert result.solve_seconds < 10


def test_e2e_generate_respects_absolute_off(scenario):
    staff_id = scenario["preference_staff"]["absolute_off"]
    assert DATES[ABSOLUTE_OFF_INDEX] not in _worked_dates(scenario["result"], staff_id)


def test_e2e_generate_respects_prefer_off_when_possible(scenario):
    result = scenario["result"]
    assert result.prefer_off_requested_total >= 2
    assert result.prefer_off_respected_total >= 1


def test_e2e_generate_uses_the_effective_times_when_scheduled(scenario):
    staff_id = scenario["preference_staff"]["time_override"]
    by_date = {
        a.work_date: a
        for a in scenario["result"].assignments
        if a.staff_id == staff_id and a.is_working
    }
    early = by_date.get(DATES[EARLY_LEAVE_INDEX])
    late = by_date.get(DATES[LATE_START_INDEX])
    if early:
        assert early.end_time == "13:00"
    if late:
        assert late.start_time == "11:00"
    # 配置の公平性を優先した結果、両日とも休みになる場合もある。
    # 実効時刻そのものは直前のconditionテストで常に検証している。


def test_e2e_generate_can_use_an_extra_available_staff(scenario):
    """通常休みでも「勤務可能」と希望した人を使って必要人数を満たせる."""
    staff_id = scenario["preference_staff"]["extra"]
    work_date = DATES[EXTRA_INDEX]
    day = next(d for d in scenario["result"].days if d.work_date == work_date)
    assert work_date in _worked_dates(scenario["result"], staff_id)
    assert day.staff_shortage == 0


def test_e2e_generate_meets_the_requirements_where_possible(scenario):
    shortage_date = DATES[SHORTAGE_INDEX]
    for day in scenario["result"].days:
        if day.work_date == shortage_date:
            continue
        assert day.staff_shortage == 0, day.work_date


def test_e2e_generate_returns_a_plan_even_with_a_shortage(scenario):
    """§38: どうしても1名足りない日があっても結果を返す."""
    day = next(d for d in scenario["result"].days if d.work_date == DATES[SHORTAGE_INDEX])
    assert day.staff_shortage == 1
    assert day.scheduled_staff_count == day.required_total_staff - 1
    assert scenario["result"].has_solution


def test_e2e_generate_meets_the_role_requirements(scenario):
    staff_roles = {
        d.staff.staff_id: d.staff.role_id for d in scenario["staff_details"]
    }
    for index in ROLE_REQUIREMENT_DAYS:
        work_date = DATES[index]
        working = [
            a.staff_id
            for a in scenario["result"].assignments
            if a.work_date == work_date and a.is_working
        ]
        roles = {staff_roles[s] for s in working}
        assert LEADER in roles, work_date
        assert CHECKER in roles, work_date


def test_e2e_generate_meets_the_skill_requirements(scenario):
    skills = {d.staff.staff_id: d.staff.skill_level for d in scenario["staff_details"]}
    for index, count in SKILL_REQUIREMENT_DAYS.items():
        work_date = DATES[index]
        high = [
            a.staff_id
            for a in scenario["result"].assignments
            if a.work_date == work_date and a.is_working and skills[a.staff_id] >= 4
        ]
        assert len(high) >= count, work_date


def test_e2e_generate_respects_max_consecutive_days(scenario):
    limits = {
        d.staff.staff_id: d.staff.max_consecutive_days
        for d in scenario["staff_details"]
        if d.staff.max_consecutive_days
    }
    for staff_id, limit in limits.items():
        worked = _worked_dates(scenario["result"], staff_id)
        run = 0
        for work_date in DATES:
            run = run + 1 if work_date in worked else 0
            assert run <= limit, (staff_id, work_date)


def test_e2e_generate_moves_towards_the_target_days(scenario):
    summaries = {s.staff_id: s for s in scenario["result"].staff_summaries}
    assert len(summaries) == 12
    with_target = [s for s in summaries.values() if s.has_target]
    assert with_target
    high = [s for s in with_target if s.target_days_per_week == 5]
    low = [s for s in with_target if s.target_days_per_week == 3]
    assert min(s.scheduled_days for s in high) >= max(s.scheduled_days for s in low)


# ---------------------------------------------------------------------------
# 5. 下書き保存
# ---------------------------------------------------------------------------


def test_e2e_save_draft(scenario):
    assert scenario["save_errors"] == []
    assert scenario["run_id"] is not None
    saved = scenario["saved"]
    assert saved.existing_dates == DATES
    assert saved.draft_dates == DATES
    assert all(len(saved.day(d).assignments) == 12 for d in DATES)


def test_e2e_saved_schedule_keeps_days_off_as_rows(scenario):
    saved = scenario["saved"]
    off_rows = [
        a
        for work_date in DATES
        for a in saved.day(work_date).assignments.values()
        if not a.is_working
    ]
    assert off_rows
    assert all(a.start_time is None and a.end_time is None for a in off_rows)


# ---------------------------------------------------------------------------
# 6. 手修正・固定
# ---------------------------------------------------------------------------


def test_e2e_manual_edit_working_to_off(scenario):
    manual = scenario["manual"]
    assert scenario["manual_errors"] == []
    row = _assignment(scenario["after_manual"], manual["to_off"], manual["staff_id"])
    assert row.is_working is False
    assert row.locked is True
    assert row.source == "MANUAL"


def test_e2e_manual_edit_off_to_working(scenario):
    manual = scenario["manual"]
    row = _assignment(scenario["after_manual"], manual["to_work"], manual["staff_id"])
    assert row.is_working is True
    assert row.locked is True
    assert row.start_time and row.end_time


def test_e2e_lock_can_be_released_explicitly(scenario):
    manual = scenario["manual"]
    before = _assignment(
        scenario["after_manual"], manual["unlock_target"], manual["staff_id"]
    )
    after = _assignment(
        scenario["after_unlock"], manual["unlock_target"], manual["staff_id"]
    )
    assert before.locked is True
    assert after.locked is False
    assert after.is_working is True


# ---------------------------------------------------------------------------
# 7. 再生成
# ---------------------------------------------------------------------------


def test_e2e_regeneration_preview_does_not_change_the_database(scenario):
    preview = scenario["preview"]
    assert preview.result is not None and preview.result.has_solution
    assert preview.conflicts == []
    before = scenario["saved"]
    unchanged = scenario["before_regeneration"]
    # previewの前後で、保存済みの勤務状態は手修正の分しか変わっていない
    manual = scenario["manual"]
    changed = {
        (work_date, staff_id)
        for work_date in DATES
        for staff_id, row in unchanged.day(work_date).assignments.items()
        if row.is_working != before.day(work_date).assignments[staff_id].is_working
    }
    assert changed <= {
        (manual["to_off"], manual["staff_id"]),
        (manual["to_work"], manual["staff_id"]),
    }


def test_e2e_regeneration_keeps_locked_cells(scenario):
    manual = scenario["manual"]
    after = scenario["after_regeneration"]
    assert scenario["apply_errors"] == []
    assert scenario["apply_run_id"] is not None

    to_off = _assignment(after, manual["to_off"], manual["staff_id"])
    to_work = _assignment(after, manual["to_work"], manual["staff_id"])
    assert (to_off.is_working, to_off.locked) == (False, True)
    assert (to_work.is_working, to_work.locked) == (True, True)


def test_e2e_regeneration_still_respects_absolute_off(scenario):
    staff_id = scenario["preference_staff"]["absolute_off"]
    row = _assignment(scenario["after_regeneration"], DATES[ABSOLUTE_OFF_INDEX], staff_id)
    assert row.is_working is False


# ---------------------------------------------------------------------------
# 8. 日別確定
# ---------------------------------------------------------------------------


def test_e2e_finalize_two_days(scenario):
    validation = scenario["validation"]
    assert validation.finalized_days == 2
    assert validation.draft_days == PERIOD_DAYS - 2
    for work_date in scenario["finalized_dates"]:
        assert validation.day(work_date).is_finalized is True


def test_e2e_finalized_days_have_no_draft_notice(scenario):
    validation = scenario["validation"]
    for work_date in scenario["finalized_dates"]:
        assert SCHEDULE_DRAFT not in [i.code for i in validation.day(work_date).issues]


# ---------------------------------------------------------------------------
# 9. 勤務表の検証
# ---------------------------------------------------------------------------


def test_e2e_validation_covers_every_day(scenario):
    validation = scenario["validation"]
    assert len(validation.days) == PERIOD_DAYS
    assert validation.existing_days == PERIOD_DAYS
    assert validation.missing_days == 0
    assert validation.requirement_missing_days == 0


def test_e2e_validation_finds_the_shortage_day(scenario):
    day = scenario["validation"].day(DATES[SHORTAGE_INDEX])
    assert day.staff_shortage == 1
    assert day.status == VALIDATION_STATUS_ERROR
    assert STAFF_SHORTAGE in [i.code for i in day.issues]


def test_e2e_validation_has_no_other_shortage(scenario):
    shortage_dates = [
        d.work_date for d in scenario["validation"].days if d.staff_shortage
    ]
    assert shortage_dates == [DATES[SHORTAGE_INDEX]]


# ---------------------------------------------------------------------------
# 10. Excel出力
# ---------------------------------------------------------------------------


def test_e2e_export_creates_a_readable_workbook(scenario):
    """§68: 実ファイルとして書き出して openpyxl で読み直せること."""
    assert scenario["excel_path"].exists()
    # BytesIOではなく書き出したファイルから読み直す（破損していないことの確認）
    reopened = load_workbook(scenario["excel_path"])
    assert reopened.sheetnames[0] == SHEET_SCHEDULE
    assert SHEET_DAILY in reopened.sheetnames
    assert SHEET_ISSUES in reopened.sheetnames


def test_e2e_export_sheet_shows_every_staff(scenario):
    ws = scenario["workbook"][SHEET_SCHEDULE]
    names = {ws.cell(row=row, column=1).value for row in range(1, ws.max_row + 1)}
    for _, name, *_ in STAFF_PLAN:
        assert name in names


def test_e2e_export_sheet_shows_times_and_days_off(scenario):
    ws = scenario["workbook"][SHEET_SCHEDULE]
    values = {
        str(cell.value)
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    }
    assert "休" in values
    assert any("09:00-15:30" == v for v in values)
    assert "確定" in values
    assert "下書き" in values


def test_e2e_export_issue_sheet_lists_the_shortage(scenario):
    ws = scenario["workbook"][SHEET_ISSUES]
    rows = [
        [ws.cell(row=row, column=col).value for col in range(1, 5)]
        for row in range(2, ws.max_row + 1)
    ]
    assert any(r[2] == STAFF_SHORTAGE and r[0] == DATES[SHORTAGE_INDEX] for r in rows)


def test_e2e_does_not_touch_attendance_data(scenario):
    """計画勤務表と勤怠実績を混ぜない."""
    conn = get_connection(str(scenario["db_path"]))
    counts = {
        table: conn.execute(f"SELECT count(*) FROM {table}").fetchone()[0]
        for table in ("attendance_imports", "attendance_shifts")
    }
    conn.close()
    assert counts == {"attendance_imports": 0, "attendance_shifts": 0}
